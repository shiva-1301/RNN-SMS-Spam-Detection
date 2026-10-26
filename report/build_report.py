"""Build report/RNN_SMS_Spam_Detection_Report.docx by editing a copy of ../Sample report.docx.

Nothing is created from scratch: every heading, paragraph, list item and table in the output
is a deep copy of an element of the sample (same fonts, sizes, numbering, spacing, table
style), with only the text replaced. The cover page, logo and index table are kept as they
are. Numbers come from outputs/*.json, so the report always matches the final run.

Usage:  python report/build_report.py [--pages "3,4,5,..."]
"""
import argparse
import copy
import json
import shutil
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT.parent / "Sample report.docx"
OUT = ROOT / "report" / "RNN_SMS_Spam_Detection_Report.docx"
OUTPUTS = ROOT / "outputs"

SAMPLE_RED = "EE0000"   # colour the sample uses to mark its example text (see NOTES.md)
TITLE = "SMS Spam Detection using Recurrent Neural Networks (RNN)"

W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


# ----------------------------------------------------------------------------- numbers
def load_numbers():
    test = json.loads((OUTPUTS / "test_metrics.json").read_text())
    results = json.loads((OUTPUTS / "experiments_results.json").read_text())
    best = json.loads((ROOT / "models" / "best_config.json").read_text())
    hist = json.loads((ROOT / "models" / "best_history.json").read_text())
    preds = json.loads((OUTPUTS / "sample_predictions.json").read_text())
    (tn, fp), (fn, tp) = test["confusion_matrix"]
    e = best["weights_from_epoch"] - 1
    return dict(test=test, results=results, best=best, hist=hist, preds=preds,
                tn=tn, fp=fp, fn=fn, tp=tp, n_test=tn + fp + fn + tp,
                acc=f"{test['accuracy'] * 100:.2f}%", prec=f"{test['precision']:.4f}",
                rec=f"{test['recall']:.4f}", f1=f"{test['f1']:.4f}",
                val_acc_best=f"{hist['val_accuracy'][e] * 100:.2f}%",
                val_loss_best=f"{hist['val_loss'][e]:.4f}",
                val_loss_last=f"{hist['val_loss'][-1]:.4f}",
                train_acc_best=f"{hist['accuracy'][e] * 100:.2f}%",
                cw_ham=f"{best['class_weight']['0']:.3f}", cw_spam=f"{best['class_weight']['1']:.3f}")


# ----------------------------------------------------------------------------- xml helpers
def strip_ids(el):
    """Drop Word paragraph ids from copies so the document has no duplicate ids."""
    for node in el.iter():
        for attr in (W14 + "paraId", W14 + "textId"):
            node.attrib.pop(attr, None)


def strip_red(rpr):
    if rpr is None:
        return
    for c in rpr.findall(qn("w:color")):
        if c.get(qn("w:val")) == SAMPLE_RED:
            rpr.remove(c)


def set_bold(rpr, bold):
    for tag in ("w:b", "w:bCs"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    if bold:
        fonts = rpr.find(qn("w:rFonts"))
        idx = list(rpr).index(fonts) + 1 if fonts is not None else 0
        for offset, tag in enumerate(("w:b", "w:bCs")):
            rpr.insert(idx + offset, rpr.makeelement(qn(tag), {}))


def first_rpr(p):
    for r in p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None and r.find(qn("w:t")) is not None:
            return rpr
    for r in p.iter(qn("w:r")):
        if r.find(qn("w:rPr")) is not None:
            return r.find(qn("w:rPr"))
    return None


def make_para(proto, runs=None, align=None):
    """Copy a sample paragraph, keep its paragraph/run formatting, replace its text.

    runs: str, or list of str / (text, bold) tuples. bold=None keeps the sample's weight."""
    p = copy.deepcopy(proto)
    strip_ids(p)
    rpr_proto = copy.deepcopy(first_rpr(proto))
    ppr = p.find(qn("w:pPr"))
    for child in list(p):
        if child is not ppr:
            p.remove(child)
    if ppr is not None:
        strip_red(ppr.find(qn("w:rPr")))
        if align:
            jc = ppr.find(qn("w:jc"))
            if jc is None:
                jc = ppr.makeelement(qn("w:jc"), {})
                ppr.append(jc)
            jc.set(qn("w:val"), align)
    if runs is None:
        return p
    if isinstance(runs, str):
        runs = [runs]
    for item in runs:
        text, bold = (item, None) if isinstance(item, str) else item
        r = p.makeelement(qn("w:r"), {})
        if rpr_proto is not None:
            rpr = copy.deepcopy(rpr_proto)
            strip_red(rpr)
            for tag in ("w:noProof", "w:lastRenderedPageBreak"):
                for el in rpr.findall(qn(tag)):
                    rpr.remove(el)
            if bold is not None:
                set_bold(rpr, bold)
            r.append(rpr)
        t = r.makeelement(qn("w:t"), {})
        t.text = text
        if text != text.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p.append(r)
    return p


def keep_with_next(p):
    """Keep a caption/label on the same page as the figure or table that follows it."""
    ppr = p.find(qn("w:pPr"))
    if ppr.find(qn("w:keepNext")) is None:
        style = ppr.find(qn("w:pStyle"))
        ppr.insert(list(ppr).index(style) + 1 if style is not None else 0,
                   ppr.makeelement(qn("w:keepNext"), {}))
    return p


def make_table(index_tbl, rows, widths, left_cols=(), indent=None):
    """Copy the sample's INDEX table (Table Grid, centred Times New Roman cells) and
    refill it with `rows` (first row = bold header). A cell may be a list of lines."""
    tbl = copy.deepcopy(index_tbl)
    strip_ids(tbl)
    trs = tbl.findall(qn("w:tr"))
    head_tc, body_tc = trs[0].find(qn("w:tc")), trs[1].find(qn("w:tc"))
    tr_proto = copy.deepcopy(trs[1])
    for tc in tr_proto.findall(qn("w:tc")):
        tr_proto.remove(tc)
    for tr in trs:
        tbl.remove(tr)
    grid = tbl.find(qn("w:tblGrid"))
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        grid.append(grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(w)}))
    if indent:
        tblpr = tbl.find(qn("w:tblPr"))
        ind = tblpr.makeelement(qn("w:tblInd"), {qn("w:w"): str(indent), qn("w:type"): "dxa"})
        tblpr.insert(list(tblpr).index(tblpr.find(qn("w:tblW"))) + 1, ind)
    for ri, row in enumerate(rows):
        tr = copy.deepcopy(tr_proto)
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = tr.makeelement(qn("w:trPr"), {})
            prex = tr.find(qn("w:tblPrEx"))
            tr.insert(list(tr).index(prex) + 1 if prex is not None else 0, trpr)
        trpr.append(trpr.makeelement(qn("w:cantSplit"), {}))
        if ri == 0:
            trpr.append(trpr.makeelement(qn("w:tblHeader"), {}))
        for ci, (cell, w) in enumerate(zip(row, widths)):
            proto_tc = head_tc if ri == 0 else body_tc
            tc = copy.deepcopy(proto_tc)
            tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(w))
            proto_p = proto_tc.find(qn("w:p"))
            for p in tc.findall(qn("w:p")):
                tc.remove(p)
            lines = cell if isinstance(cell, list) else [cell]
            align = "left" if (ci in left_cols and ri > 0) else None
            for line in lines:
                tc.append(make_para(proto_p, line, align=align))
            tr.append(tc)
        tbl.append(tr)
    return tbl


# ----------------------------------------------------------------------------- content
def build(page_numbers=None):
    n = load_numbers()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SAMPLE, OUT)
    doc = docx.Document(OUT)
    P = [p._p for p in doc.paragraphs]
    index_tbl = doc.tables[0]._tbl
    body = doc.element.body

    def proto(i):
        return copy.deepcopy(P[i])

    H = {name: proto(i) for name, i in dict(
        title=21, abstract=22, intro=48, objectives=58, dataset=70, model=96, framework=99,
        method=119, performance=143, graphs=202, analysis=252, conclusion=265).items()}
    T = {name: proto(i) for name, i in dict(
        body_red=50, body_black=121, sub_bold=120, obj_bullet=60, ds_intro=73, ds_bullet=74,
        ds_final=77, fw_item=101, fw_desc=102, mp_head=148, mp_quote=149, mp_text=150,
        mp_bullet=152, program=157, mp_numhead=160, mp_image=158, g_quote=204, g_text=206,
        g_bullet=208, g_label=218, g_image=219, a_text=254, a_bullet=258, a_close=263,
        c_text=268, page_break=66).items()}

    # ---- cover page: only the title line changes (placeholders stay as in the sample)
    cover = doc.paragraphs[0]
    cover.runs[0].text = TITLE
    for r in cover.runs[1:]:
        r._r.getparent().remove(r._r)

    # ---- drop everything from the old "TITLE" page to the end (sectPr stays)
    start = list(body).index(P[21])
    for el in list(body)[start:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    sect = body.find(qn("w:sectPr"))

    def add(el):
        sect.addprevious(el)
        return el

    def page_break():
        add(copy.deepcopy(T["page_break"]))

    def heading(name, text=None):
        h = copy.deepcopy(H[name])
        strip_ids(h)
        has_break = any(br.get(qn("w:type")) == "page" for br in h.iter(qn("w:br")))
        if name not in ("title", "abstract") and not has_break:
            page_break()
        if text:
            runs = h.findall(qn("w:r"))
            text_runs = [r for r in runs if r.find(qn("w:t")) is not None]
            text_runs[0].find(qn("w:t")).text = text
            for r in text_runs[1:]:
                h.remove(r)
        add(h)

    def para(kind, runs, **kw):
        p = make_para(T[kind], runs, **kw)
        if kind in ("program", "g_label"):
            keep_with_next(p)
        return add(p)

    def blank(kind="body_black"):
        add(make_para(T[kind], None))

    def table(rows, widths, left_cols=(), indent=None):
        add(make_table(index_tbl, rows, widths, left_cols, indent))

    def image(kind, path, width_in):
        p = add(make_para(T[kind], None))
        Paragraph(p, doc._body).add_run().add_picture(str(path), width=Inches(width_in))

    res = n["results"]
    best_name = n["best"]["name"]

    # ---- TITLE + ABSTRACT -------------------------------------------------------------
    heading("title", TITLE.upper())
    heading("abstract")
    para("body_red", [
        "Unwanted SMS spam such as fake prize offers, premium-rate numbers and phishing links is "
        "a daily nuisance and a security risk for mobile users. This project presents an ",
        ("“RNN-based SMS Spam Detection System”", True), " developed in ",
        ("Python", True), " using ", ("TensorFlow/Keras", True),
        ". Messages from the UCI SMS Spam Collection (5,574 real SMS messages, 13.4% spam) are "
        "lowercased, cleaned, tokenized and padded to 50 words, and then classified by a ",
        ("Recurrent Neural Network", True),
        " built as Embedding → recurrent layer → Dense → Sigmoid. Class weights handle "
        "the imbalance between ham and spam, and dropout with early stopping controls overfitting. "
        "Six configurations (SimpleRNN, LSTM and GRU) were trained and the best one was chosen on "
        f"the validation set. The final GRU-based model achieved {n['acc']} accuracy, "
        f"{n['prec']} precision, {n['rec']} recall and an F1-score of {n['f1']} for the spam class "
        f"on {n['n_test']} unseen test messages, and a prediction script labels any new SMS as "
        "spam or ham instantly."])

    # ---- INTRODUCTION -------------------------------------------------------------------
    heading("intro")
    para("body_red", "SMS is still one of the most widely used ways to communicate, which also "
         "makes it an easy target for spammers. Spam messages advertise fake prizes, push "
         "premium-rate numbers or carry phishing links that steal personal and banking details, "
         "so detecting them automatically is an important and practical problem.")
    para("body_red", ["Traditional rule-based filters depend on fixed keyword lists and are easy "
         "to bypass when spammers change their wording. This project addresses that gap by using a ",
         ("Recurrent Neural Network (RNN)", True), ", which reads a message word by word and "
         "learns from the order and context of the words instead of isolated keywords."])
    para("body_red", "The system is designed to help mobile users and service providers separate "
         "spam from legitimate (ham) messages automatically, by learning directly from thousands of "
         "real, labelled SMS messages rather than from hand-written rules.")

    # ---- OBJECTIVES ---------------------------------------------------------------------
    heading("objectives")
    for item in [
        ["To build a deep learning model using a ", ("Recurrent Neural Network", True),
         " that classifies SMS messages as spam or ham with high accuracy."],
        "To preprocess raw SMS text (cleaning, tokenization and padding) into a form a neural "
        "network can learn from.",
        "To handle the imbalance between ham and spam messages with class weights, and to control "
        "overfitting with dropout and early stopping.",
        "To evaluate the model on unseen data using accuracy, precision, recall, F1-score and a "
        "confusion matrix.",
        "To provide a simple prediction script that labels any custom SMS as spam or ham.",
    ]:
        para("obj_bullet", item)

    # ---- DATASET DESCRIPTION ------------------------------------------------------------
    heading("dataset")
    blank("ds_intro")
    para("ds_intro", "The project uses the SMS Spam Collection dataset from the UCI Machine "
         "Learning Repository: a public set of real English SMS messages, each tagged as ham "
         "(legitimate) or spam.")
    for label, text in [
        ("Total Records: ", "5,574 SMS messages."),
        ("Classes: ", "ham (legitimate) – 4,827 messages (86.6%); spam – 747 messages "
         "(13.4%). The dataset is imbalanced."),
        ("Source: ", "UCI Machine Learning Repository, SMS Spam Collection v.1 (Almeida, Gómez "
         "Hidalgo and Yamakami, 2011), downloaded automatically by the project."),
        ("Features: ", "one raw text field (the SMS) and one label (ham/spam). After cleaning, a "
         "message has 16.5 words on average (median 13), and 98.7% have 50 words or fewer."),
        ("Sample text: ", "spam – “Free entry in 2 a wkly comp to win FA Cup final tkts "
         "21st May 2005. Text FA to 87121 to receive entry question…”; ham – "
         "“Ok lar... Joking wif u oni...”"),
        ("Data Split: ", "stratified 70/15/15 – 3,901 training, 836 validation and 837 test "
         "messages, each keeping about 13.4% spam."),
    ]:
        para("ds_bullet", [(label, True), (text, False)])
    para("ds_final", "The tokenizer was fitted on the training messages only, so no information "
         "from the validation or test messages leaked into training.")

    # ---- MODEL DESCRIPTION --------------------------------------------------------------
    heading("model")
    cfg = n["best"]["config"]["model"]
    para("sub_bold", "1. Architecture")
    para("body_black", "The final model is a Recurrent Neural Network built with TensorFlow/Keras "
         "as a simple stack: Embedding → GRU (recurrent layer) → Dense (ReLU) → "
         "Dropout → Dense (Sigmoid). Six variants were trained (SimpleRNN, LSTM and GRU), and "
         "the GRU variant was selected because it gave the best validation performance.")
    layer_rows = [
        ["Layer", "Output Shape", "Parameters"],
        ["Input (word indices)", "(50)", "0"],
        ["Embedding (5,000 words × 64)", "(50, 64)", "320,000"],
        ["GRU (64 units, padding masked)", "(64)", "24,960"],
        ["Dense (32, ReLU)", "(32)", "2,080"],
        ["Dropout (0.3)", "(32)", "0"],
        ["Dense (1, Sigmoid) – output", "(1)", "33"],
        ["Total trainable parameters", "", f"{n['best']['params']:,}"],
    ]
    assert 320000 + 24960 + 2080 + 33 == n["best"]["params"]
    assert cfg["rnn_type"] == "gru" and cfg["embed_dim"] == 64 and cfg["rnn_units"] == 64
    table(layer_rows, [4800, 2000, 2216])
    blank()
    para("body_black", "The Embedding layer turns each word index into a 64-value vector. The GRU "
         "reads the 50 word vectors in order and keeps a 64-value hidden state that summarises the "
         "message so far; its update and reset gates decide what to keep and what to forget, which "
         "makes it more stable than a plain SimpleRNN. The Dense layer combines these features and "
         "the final Sigmoid neuron outputs the probability that the message is spam (0.5 or more "
         "is labelled spam).")
    para("sub_bold", "2. Training Process")
    para("body_black", "The model was trained with the Adam optimizer and binary cross-entropy "
         "loss in mini-batches of 32 messages. Class weights (ham "
         f"{n['cw_ham']}, spam {n['cw_spam']}) make each spam example count more, compensating for "
         "the small number of spam messages. Early stopping watched the validation loss with a "
         f"patience of 5 epochs and restored the best weights: training ran for "
         f"{n['best']['epochs_run']} epochs and the weights from epoch "
         f"{n['best']['weights_from_epoch']} were kept.")
    para("sub_bold", "3. Hyperparameters")
    table([
        ["Hyperparameter", "Value", "Reason"],
        ["Vocabulary size", "5,000 words", "Most frequent training words"],
        ["Sequence length", "50 tokens", "Covers 98.7% of messages"],
        ["Embedding dimension", "64", "Learned word vectors"],
        ["Recurrent layer", "GRU, 64 units", "Best of 6 attempts (validation)"],
        ["Dense layer", "32 units, ReLU", "Combines the GRU features"],
        ["Dropout", "0.3", "Reduces overfitting"],
        ["Optimizer", "Adam, lr 0.001, clipnorm 1.0", "Stable gradient updates"],
        ["Loss function", "Binary cross-entropy", "Two classes (spam / ham)"],
        ["Batch size", "32", "Standard mini-batch size"],
        ["Epochs", "Max 40, early stopping (patience 5)", "Stops when val. loss stops improving"],
        ["Class weights", f"ham {n['cw_ham']}, spam {n['cw_spam']}", "Balanced, spam weight halved"],
        ["Decision threshold", "0.5", "Probability ≥ 0.5 → SPAM"],
    ], [2300, 3500, 3216], left_cols=(0, 1, 2))

    # ---- FRAMEWORK AND TOOLS USED -------------------------------------------------------
    heading("framework")
    for item, desc in [
        ("Python 3", "Core programming language for data processing, model building and scripts."),
        ("TensorFlow / Keras 2.21", "Deep learning framework used to build, train and save the RNN model."),
        ("scikit-learn", "Used for the stratified data split, class weights and evaluation metrics."),
        ("pandas & NumPy", "Used to load the dataset and handle the padded sequences as arrays."),
        ("Matplotlib", "Used to plot the accuracy graph, loss graph and confusion matrix."),
        ("Docker", "Provides a reproducible Linux environment with TensorFlow for training."),
        ("UCI SMS Spam Collection", "The dataset of 5,574 labelled SMS messages."),
    ]:
        para("fw_item", item)
        para("fw_desc", desc)

    # ---- METHODOLOGY --------------------------------------------------------------------
    heading("method")
    for title, text in [
        ("1. Data Collection and Cleaning", "The dataset was downloaded automatically from the UCI "
         "repository. Every message was converted to lowercase; web links, e-mail addresses, "
         "currency symbols and numbers were replaced by placeholder words (urltoken, emailtoken, "
         "currencytoken, numtoken) because they are strong spam signals, and all remaining "
         "punctuation was removed."),
        ("2. Tokenization and Padding", "A word-level tokenizer was built from the training "
         "messages, keeping the 5,000 most frequent words (unknown words map to a special "
         "out-of-vocabulary index). Each message was converted into a sequence of word indices and "
         "padded with zeros at the front, or truncated, to exactly 50 tokens."),
        ("3. Data Splitting and Class Imbalance", "The data was split with stratification into "
         "70% training, 15% validation and 15% test sets, so every set keeps the same spam ratio. "
         "Because only 13.4% of the messages are spam, class weights were applied during training "
         "so that missing a spam message costs more than misclassifying a ham message."),
        ("4. Model Training and Selection", "A SimpleRNN baseline was trained first and then "
         "improved over five further attempts (dropout, early stopping, a smaller learning rate "
         "with gradient clipping, LSTM, GRU and softer class weights). The model with the best "
         "validation F1-score was kept, and only then was it evaluated once on the test set."),
        ("5. Evaluation and Prediction", f"Accuracy, precision, recall, F1-score and the confusion "
         f"matrix were computed on the {n['n_test']} test messages. A prediction script applies "
         "the same cleaning and tokenization to any new SMS and prints SPAM or HAM with its "
         "probability."),
    ]:
        para("sub_bold", title)
        para("body_black", text)

    # ---- MODEL PERFORMANCE --------------------------------------------------------------
    heading("performance")
    blank()
    para("mp_head", "1. DATA PREPARATION & CLEANING")
    para("mp_quote", "“Convert raw SMS text into fixed-length sequences of numbers.”")
    para("mp_text", "Raw messages contain capital letters, punctuation, links and phone numbers. "
         "The cleaning function normalises them, the tokenizer maps every word to an index and the "
         "sequences are padded to the same length:")
    for item in ["Lowercasing and punctuation removal",
                 "Links, e-mails, currency symbols and numbers → placeholder words",
                 "5,000 most frequent training words kept; other words → OOV index",
                 "Pre-padding / truncation to 50 tokens"]:
        para("mp_bullet", item)
    para("program", "Program:")
    table([["src/data_prep.py (excerpt)"], [[
        "def clean_text(text):",
        "    text = text.lower()",
        "    text = URL.sub(\" urltoken \", text)",
        "    text = EMAIL.sub(\" emailtoken \", text)",
        "    text = CURRENCY.sub(\" currencytoken \", text)",
        "    text = NUMBER.sub(\" numtoken \", text)",
        "    text = NON_ALNUM.sub(\" \", text)",
        "    return SPACES.sub(\" \", text).strip()",
        "",
        "tokenizer = SimpleTokenizer(vocab_size=5000).fit(train_df[\"clean\"])",
        "X_train = pad(tokenizer.texts_to_sequences(train_df[\"clean\"]))  # 50 tokens",
    ]]], [9016], left_cols=(0,))
    blank()
    para("mp_numhead", "THE RECURRENT NEURAL NETWORK (RNN)")
    para("mp_quote", "“Train the RNN to learn the patterns of spam messages.”")
    para("mp_text", "The model reads each message word by word. At every step the recurrent layer "
         "updates its hidden state, so after the last word it holds a summary of the whole "
         "message, which the Dense and Sigmoid layers turn into a spam probability. Six "
         "configurations were trained in the self-evaluation loop and compared on the validation "
         "set:")
    rows = [["No.", "Model and key changes", "Val F1", "Test Acc.", "Test F1"]]
    labels = {
        "A1_baseline": "SimpleRNN baseline (no dropout, 15 epochs)",
        "A2_simplernn_dropout_es": "SimpleRNN + dropout 0.3 + early stopping (patience 3)",
        "A3_simplernn_small_lr": "Smaller SimpleRNN, lr 0.0005, gradient clipping, patience 5",
        "A4_lstm": "LSTM, dropout 0.3, early stopping",
        "A5_gru": "GRU, dropout 0.3, early stopping",
        "A6_gru_softer_weights": "GRU + softer spam class weight",
    }
    for i, (name, r) in enumerate(res.items(), 1):
        label = labels[name] + (" (selected)" if name == best_name else "")
        rows.append([str(i), label, f"{r['val']['f1']:.4f}", f"{r['test']['accuracy'] * 100:.2f}%",
                     f"{r['test']['f1']:.4f}"])
    table(rows, [700, 4916, 1100, 1200, 1100], left_cols=(1,))
    para("mp_text", "The baseline already reached the targets but overfitted (100% training "
         "accuracy while its validation loss rose), so regularisation was added. The final model "
         "was chosen by validation F1-score, not by test score.")
    para("program", "Program:")
    table([["src/model.py and src/train.py (excerpt)"], [[
        "model = tf.keras.Sequential([",
        "    layers.Input(shape=(50,), dtype=\"int32\"),",
        "    layers.Embedding(5000, 64, mask_zero=True),",
        "    layers.GRU(64, dropout=0.3),",
        "    layers.Dense(32, activation=\"relu\"),",
        "    layers.Dropout(0.3),",
        "    layers.Dense(1, activation=\"sigmoid\")])",
        "model.compile(optimizer=Adam(0.001, clipnorm=1.0),",
        "              loss=\"binary_crossentropy\", metrics=[\"accuracy\"])",
        "",
        "w = compute_class_weight(\"balanced\", classes=[0, 1], y=y_train)",
        "class_weight = {0: w[0], 1: w[1] * 0.5}",
        "early_stop = EarlyStopping(monitor=\"val_loss\", patience=5,",
        "                           restore_best_weights=True)",
        "history = model.fit(X_train, y_train, validation_data=(X_val, y_val),",
        "                    epochs=40, batch_size=32, class_weight=class_weight,",
        "                    callbacks=[early_stop])",
        "model.save(\"models/best_model.keras\")",
    ]]], [9016], left_cols=(0,))
    blank()
    para("mp_numhead", "PREDICTION ON NEW MESSAGES")
    para("mp_quote", "“Classify any custom SMS as SPAM or HAM.”")
    para("mp_text", "The prediction script loads the saved model and tokenizer, applies the same "
         "cleaning and padding, and prints the label with the spam probability. It was tested on "
         "eight new messages written for this check (none are from the dataset):")
    para("program", "Program:")
    table([["src/predict.py (usage)"], [[
        "python src/predict.py \"WINNER!! Claim your free prize now, call 09061701461\"",
        "python src/predict.py          # interactive mode",
    ]]], [9016], left_cols=(0,))
    blank()
    para("program", "Output:")
    prows = [["Message", "Expected", "Prediction (spam probability)"]]
    for r in n["preds"]:
        prows.append([r["message"], r["expected"], f"{r['predicted']} ({r['spam_probability']:.4f})"])
    table(prows, [5800, 1200, 2016], left_cols=(0,))
    blank()
    correct = sum(r["correct"] for r in n["preds"])
    para("mp_text", f"All {correct} of {len(n['preds'])} messages were classified correctly, "
         "including a phishing message that contains a link but no “prize” words.")

    # ---- GRAPHS -------------------------------------------------------------------------
    heading("graphs")
    blank("g_text")
    para("g_quote", "“To verify training, the accuracy and loss on the training and validation "
         "sets are plotted for every epoch, and the test predictions are summarised in a "
         "confusion matrix.”")
    blank("g_text")
    para("g_text", ["The graphs are produced from the saved training history and test predictions "
         "of the final model (", ("GRU", True), ", attempt 6):"])
    blank("g_text")
    for item in [
        "The blue line shows the training set and the orange line the validation set.",
        f"Validation accuracy reaches {n['val_acc_best']} at epoch {n['best']['weights_from_epoch']}, "
        "the epoch whose weights were restored (dashed line).",
        f"After epoch {n['best']['weights_from_epoch']} the training loss keeps falling towards zero "
        f"while the validation loss rises from {n['val_loss_best']} to {n['val_loss_last']} – the "
        "start of overfitting, which early stopping prevents.",
        f"The confusion matrix shows only {n['fp']} false positive and {n['fn']} false negatives "
        f"among {n['n_test']} test messages.",
    ]:
        para("g_bullet", item)
    blank("g_text")
    para("g_label", "Figure 1: Training vs Validation Accuracy")
    image("g_image", OUTPUTS / "accuracy_graph.png", 5.6)
    blank("g_text")
    para("g_label", "Figure 2: Training vs Validation Loss")
    image("g_image", OUTPUTS / "loss_graph.png", 5.6)
    blank("g_text")
    para("g_label", "Figure 3: Confusion Matrix on the Test Set")
    image("g_image", OUTPUTS / "confusion_matrix.png", 4.2)

    # ---- ANALYSIS & RESULTS -------------------------------------------------------------
    heading("analysis")
    blank("a_text")
    para("a_text", ["The ", ("“RNN SMS Spam Detection”", True), " model was evaluated on "
         f"{n['n_test']} test messages ({n['tn'] + n['fp']} ham, {n['fn'] + n['tp']} spam) that were "
         "never used for training or for choosing the model."])
    blank("a_text")
    table([
        ["Metric", "Value", "Meaning"],
        ["Accuracy", n["acc"], f"{n['tn'] + n['tp']} of {n['n_test']} messages correct"],
        ["Precision (spam)", n["prec"], "Share of flagged messages that are really spam"],
        ["Recall (spam)", n["rec"], "Share of spam messages that were caught"],
        ["F1-score (spam)", n["f1"], "Balance of precision and recall"],
        ["Confusion matrix", f"TN {n['tn']}, FP {n['fp']}, FN {n['fn']}, TP {n['tp']}",
         "Ham/spam counts, actual vs predicted"],
    ], [2000, 3000, 3296], left_cols=(2,), indent=720)
    blank("a_text")
    para("a_text", "The analysis revealed:")
    blank("a_text")
    a3, a4, a5 = res["A3_simplernn_small_lr"], res["A4_lstm"], res["A5_gru"]
    for item in [
        f"Accuracy: The final GRU model classified {n['acc']} of the test messages correctly, above "
        "the 95% target.",
        f"Precision: {n['prec']} – of the {n['tp'] + n['fp']} messages flagged as spam, "
        f"{n['tp']} really were spam; only {n['fp']} legitimate message was wrongly flagged.",
        f"Recall: {n['rec']} – {n['tp']} of the {n['tp'] + n['fn']} spam messages were caught "
        f"and {n['fn']} were missed.",
        f"F1-score: {n['f1']} for the spam class, above the 0.90 target.",
        "Model Comparison: after tuning, the SimpleRNN reached "
        f"{a3['test']['accuracy'] * 100:.2f}% accuracy (F1 {a3['test']['f1']:.4f}), the LSTM "
        f"{a4['test']['accuracy'] * 100:.2f}% (F1 {a4['test']['f1']:.4f}) and the GRU "
        f"{n['test']['accuracy'] * 100:.2f}–{a5['test']['accuracy'] * 100:.2f}% (F1 "
        f"{n['test']['f1']:.4f}–{a5['test']['f1']:.4f}); the gated cells trained more stably.",
        "Overfitting Control: without regularisation the baseline reached 100% training accuracy "
        "while its validation loss rose; dropout and early stopping (best weights from epoch "
        f"{n['best']['weights_from_epoch']}) closed this gap.",
        f"Custom Messages: the prediction script labelled all {len(n['preds'])} new test messages "
        "correctly (4 spam, 4 ham).",
    ]:
        label, text = item.split(":", 1)
        para("a_bullet", [(label + ":", True), (text, False)])
    para("a_close", "These findings show that a compact recurrent neural network can filter SMS "
         "spam very accurately, catching most spam while blocking almost no legitimate messages.")

    # ---- CONCLUSION ---------------------------------------------------------------------
    heading("conclusion")
    blank("c_text")
    para("c_text", ["The ", ("RNN SMS Spam Detection", True), " project successfully demonstrates how "
         "deep learning can automatically separate spam from legitimate text messages. Using the "
         "real UCI SMS Spam Collection, a simple pipeline of text cleaning, tokenization, padding "
         "and a recurrent neural network (Embedding → GRU → Dense → Sigmoid) achieved "
         f"{n['acc']} accuracy and a spam F1-score of {n['f1']} on unseen test messages, with only "
         f"{n['fp']} legitimate message misclassified."])
    para("c_text", "The study confirms that handling class imbalance and controlling overfitting "
         "matter as much as the choice of model: class weights, dropout and early stopping turned "
         "an overfitting baseline into a stable model that was selected purely on validation data. "
         "In future, the system can be extended with a larger and more recent multilingual SMS "
         "dataset, pre-trained word embeddings or attention-based models, and deployed as a mobile "
         "or web service for real-time filtering.")

    # ---- INDEX page numbers (filled after a render pass) --------------------------------
    if page_numbers:
        rows = doc.tables[0].rows[1:]
        assert len(rows) == len(page_numbers)
        for row, num in zip(rows, page_numbers):
            cell_p = row.cells[2].paragraphs[0]._p
            sno_p = row.cells[0].paragraphs[0]._p
            new_p = make_para(sno_p, str(num))
            cell_p.addprevious(new_p)
            cell_p.getparent().remove(cell_p)

    # ---- remove the sample's old screenshots (no longer referenced) ---------------------
    xml = doc.element.xml
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.IMAGE and f'"{rid}"' not in xml:
            doc.part.drop_rel(rid)

    doc.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--pages", help="comma-separated page numbers for the INDEX table")
    args = ap.parse_args()
    build([int(x) for x in args.pages.split(",")] if args.pages else None)
